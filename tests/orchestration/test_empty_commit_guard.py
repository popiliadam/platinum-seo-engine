"""V-01: an empty or sharply shrunk transform output must NEVER replace a sheet.

``transaction.replace`` clears the data block and then writes ``rows``, so a
0-row commit silently empties the sheet. The count checks used to run AFTER the
write (run_step) or not at all (empty input, model_attested steps). These tests
pin the order: count first, write only if the count is acceptable.
"""
from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from scripts.orchestration import committer, workflow_driver
from scripts.orchestration.committer import EmptyCommitError
from scripts.orchestration.run_step import StepSpec, run_step

RUN_ID = "test-proj-2026-10-19-v01a"
SLUG = "test-proj"
NOW = 1_750_000_000
SHEET = "topical_map"


def _iso(epoch: float) -> str:
    return datetime.fromtimestamp(epoch, tz=timezone.utc).isoformat()


def _drop(path: Path, rows: list[dict]) -> Path:
    payload = {
        "provenance": {
            "run_id": RUN_ID,
            "slug": SLUG,
            "fetched_at": _iso(NOW - 60),
            "declared_count": len(rows),
        },
        "rows": rows,
    }
    path.write_text(json.dumps(payload), encoding="utf-8")
    os.utime(path, (NOW - 60, NOW - 60))
    return path


class _RecordingCommit:
    """Fake commit_fn that records every call (and its kwargs)."""

    def __init__(self) -> None:
        self.calls: list[dict] = []

    def __call__(self, workbook_path, sheet, rows, **kwargs):
        self.calls.append({"sheet": sheet, "rows": list(rows), **kwargs})
        return SimpleNamespace(rows_affected=len(rows))


def _rows(n: int) -> list[dict]:
    return [{"q": f"kw{i}"} for i in range(n)]


def _spec(raw: Path, out_rows: list[dict], *, verification_class="code_verified",
          allow_empty=False) -> StepSpec:
    return StepSpec(
        name="step", raw_path=raw, sheet=SHEET,
        transform=lambda _raw: list(out_rows),
        verification_class=verification_class, allow_empty=allow_empty,
    )


def _run(spec: StepSpec, tmp_path: Path, commit: _RecordingCommit) -> dict:
    return run_step(
        spec, run_id=RUN_ID, project_slug=SLUG, workspace_root=tmp_path,
        workbook_path=tmp_path / "master.xlsx", now_epoch=NOW, commit_fn=commit,
    )


def _run_attested(spec: StepSpec, tmp_path: Path, commit: _RecordingCommit) -> dict:
    return workflow_driver._run_attested_step(
        spec, run_id=RUN_ID, project_slug=SLUG,
        workbook_path=tmp_path / "master.xlsx", now_epoch=NOW, commit_fn=commit,
    )


# --- run_step (code_verified) ------------------------------------------------

def test_empty_output_from_full_input_fails_without_commit(tmp_path: Path) -> None:
    commit = _RecordingCommit()
    step = _run(_spec(_drop(tmp_path / "d.json", _rows(50)), []), tmp_path, commit)
    assert step["status"] == "failed"
    assert commit.calls == []
    assert step["input_count"] == 50
    assert step["scored_count"] == 0


def test_empty_input_and_output_fails_without_commit(tmp_path: Path) -> None:
    """GSC returned 0 rows: used to be 'satisfied' and wipe gsc_performance."""
    commit = _RecordingCommit()
    step = _run(_spec(_drop(tmp_path / "d.json", []), []), tmp_path, commit)
    assert step["status"] == "failed"
    assert commit.calls == []


def test_sharp_drop_is_rejected_before_write(tmp_path: Path) -> None:
    commit = _RecordingCommit()
    step = _run(_spec(_drop(tmp_path / "d.json", _rows(1000)), _rows(10)), tmp_path, commit)
    assert step["status"] == "failed"
    assert commit.calls == []  # used to write 10 rows over 1000, THEN say failed


def test_half_kept_still_commits(tmp_path: Path) -> None:
    commit = _RecordingCommit()
    step = _run(_spec(_drop(tmp_path / "d.json", _rows(10)), _rows(5)), tmp_path, commit)
    assert step["status"] == "satisfied"
    assert len(commit.calls) == 1


def test_allow_empty_step_commits_empty(tmp_path: Path) -> None:
    commit = _RecordingCommit()
    spec = _spec(_drop(tmp_path / "d.json", []), [], allow_empty=True)
    step = _run(spec, tmp_path, commit)
    assert step["status"] == "satisfied"
    assert len(commit.calls) == 1
    assert commit.calls[0]["allow_empty"] is True


# --- model_attested ----------------------------------------------------------

def test_attested_empty_output_fails_without_commit(tmp_path: Path) -> None:
    commit = _RecordingCommit()
    spec = _spec(_drop(tmp_path / "d.json", _rows(40)), [],
                 verification_class="model_attested")
    step = _run_attested(spec, tmp_path, commit)
    assert step["status"] == "failed"
    assert commit.calls == []


def test_attested_aggregation_still_commits(tmp_path: Path) -> None:
    """Aggregating 40 findings into 3 rows is legitimate for an attested step."""
    commit = _RecordingCommit()
    spec = _spec(_drop(tmp_path / "d.json", _rows(40)), _rows(3),
                 verification_class="model_attested")
    step = _run_attested(spec, tmp_path, commit)
    assert step["status"] == "satisfied"
    assert len(commit.calls) == 1


def test_build_steps_carries_allow_empty(tmp_path: Path) -> None:
    table = (
        {"name": "a", "sheet": SHEET, "output_file": "a.json",
         "verification_class": "code_verified", "site_url": False, "tool": None},
        {"name": "b", "sheet": SHEET, "output_file": "b.json",
         "verification_class": "code_verified", "site_url": False, "tool": None,
         "allow_empty": True},
    )
    specs = workflow_driver.build_steps(table, RUN_ID, SLUG, tmp_path)
    assert [s.allow_empty for s in specs] == [False, True]


# --- committer (the one commit path; real transaction.replace) --------------

def _topical_row(i: int) -> dict:
    return {
        "pillar": "P01", "cluster": "c", "primary_keyword": f"kw {i}",
        "monthly_volume": 10, "data_source": "gsc", "assigned_url": f"/u{i}",
        "page_type": "pillar", "status": "TODO", "priority": "HIGH", "note": "",
    }


def _workbook(tmp_path: Path) -> Path:
    proj = tmp_path / "projects" / SLUG
    (proj / "_state").mkdir(parents=True, exist_ok=True)
    return proj / "master.xlsx"


def test_committer_refuses_empty_rows_and_keeps_sheet(tmp_path: Path) -> None:
    wb = _workbook(tmp_path)
    committer.commit(wb, SHEET, [_topical_row(i) for i in range(3)],
                     run_id=RUN_ID, project_slug=SLUG)
    with pytest.raises(EmptyCommitError):
        committer.commit(wb, SHEET, [], run_id=RUN_ID, project_slug=SLUG)
    ws = load_workbook(wb)[SHEET]
    assert ws.max_row == 7  # the 3 data rows (5..7) survived
    assert ws["C5"].value == "kw 0"


def test_committer_allow_empty_clears(tmp_path: Path) -> None:
    wb = _workbook(tmp_path)
    committer.commit(wb, SHEET, [_topical_row(0)], run_id=RUN_ID, project_slug=SLUG)
    result = committer.commit(wb, SHEET, [], run_id=RUN_ID, project_slug=SLUG,
                              allow_empty=True)
    assert result.rows_affected == 0

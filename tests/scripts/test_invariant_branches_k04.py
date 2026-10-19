"""K-02 / K-04 — invariant branches that were never exercised.

F-09 FAIL (duplicate task_id), F-10 FAIL (non-idempotent quick_wins URL) and
F-15 PASS had no test; F-05 compared only the header COUNT, so a renamed or
swapped column passed. D-03 is new: every URL in the engine's join sheets must
be normalize_url-idempotent (routed to manual triage, never blocking a commit).
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from scripts.validation import validate_invariants as vi

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "master-excel.xlsx"
SCHEMA = json.loads((ROOT / "schemas" / "master-excel.schema.json").read_text(encoding="utf-8"))


def _wb():
    return load_workbook(TEMPLATE)  # in memory; nothing is saved


def _put(wb, sheet: str, values: dict, offset: int = 1) -> None:
    ws = wb[sheet]
    hr = SCHEMA["sheets"][sheet]["header_row"]
    headers = [c.value for c in ws[hr]]
    for name, value in values.items():
        ws.cell(row=hr + offset, column=headers.index(name) + 1, value=value)


def test_f09_duplicate_task_id_fails() -> None:
    wb = _wb()
    _put(wb, "master_task", {"task_id": "T-00001"}, 1)
    _put(wb, "master_task", {"task_id": "T-00001"}, 2)
    assert vi.check_F_09(wb, "p")["verdict"] == "FAIL"


def test_f10_raw_quick_wins_url_fails() -> None:
    wb = _wb()
    _put(wb, "quick_wins", {"url": "https://X.test/a/"})
    res = vi.check_F_10(wb, "p")
    assert res["verdict"] == "FAIL" and "X.test" in res["sample_violations"][0]


def test_f15_empty_cannibalization_passes() -> None:
    assert vi.check_F_15(_wb(), "p")["verdict"] == "PASS"


def test_f05_swapped_column_names_fail() -> None:
    wb = _wb()
    ws = wb["gsc_performance"]
    hr = SCHEMA["sheets"]["gsc_performance"]["header_row"]
    a, b = ws.cell(row=hr, column=2).value, ws.cell(row=hr, column=3).value
    ws.cell(row=hr, column=2, value=b)
    ws.cell(row=hr, column=3, value=a)
    res = vi.check_F_05(wb, "p")
    assert res["verdict"] == "FAIL"
    assert any("gsc_performance: column 2" in v for v in res["sample_violations"])


def test_d03_raw_url_in_join_sheet_fails_but_only_needs_triage() -> None:
    wb = _wb()
    _put(wb, "gsc_performance", {"url": "https://x.test/kombi/"})
    res = vi.check_D_03(wb, "p")
    assert res["verdict"] == "FAIL" and res["manual_triage"] is True
    agg = vi.aggregate_verdicts([res])
    assert agg["overall"] == "AMBER" and "D-03" in agg["manual_review_required"]


def test_d03_hand_curated_display_urls_are_out_of_scope() -> None:
    wb = _wb()
    _put(wb, "topical_map", {"assigned_url": "https://x.test/kombi/"})
    assert vi.check_D_03(wb, "p")["verdict"] == "PASS"

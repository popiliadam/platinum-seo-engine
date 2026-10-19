"""V-03 / V-04: read-merge-write of a shared sheet happens under ONE excel.lock.

sheet_merge used to read the sheet without the lock, merge in memory, and lock
only for the final replace — two skills reading the same snapshot silently
dropped one another's rows. sf_import replaced robots_txt / redirect_404 with
only its own rows, wiping HF-/FN-/RP- findings and migration-map redirects.
"""
from __future__ import annotations

import multiprocessing as mp
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.excel import transaction
from scripts.excel.transaction import EmptyMergeError, LockHeldError
from scripts.ingestion import sf_import, sf_projection
from scripts.orchestration import committer
from scripts.util import sheet_merge

SLUG = "test-proj"
RUN_ID = "test-proj-2026-10-19-v03a"


def _wb(tmp_path: Path) -> Path:
    proj = tmp_path / "projects" / SLUG
    (proj / "_state").mkdir(parents=True, exist_ok=True)
    return proj / "master.xlsx"


def _robots(id_: str) -> dict:
    return {"id": id_, "level": "MEDIUM", "issue": "i", "detail": "d", "resolution": "r"}


def _redirect(url: str, action: str) -> dict:
    return {"url": url, "inlinks": 1, "action": action, "target_url": "/t", "status": "TODO"}


def _column(wb_path: Path, sheet: str, col: int = 0) -> list:
    ws = load_workbook(wb_path)[sheet]
    return [r[col] for r in ws.iter_rows(min_row=5, values_only=True) if r and r[col] is not None]


# --- transaction.merge -------------------------------------------------------

def test_merge_runs_fn_while_holding_the_lock(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    committer.commit(wb, "robots_txt", [_robots("R-001")], run_id=RUN_ID, project_slug=SLUG)
    seen: list[list[str]] = []

    def fn(existing: list[dict]) -> list[dict]:
        seen.append([r["id"] for r in existing])
        with pytest.raises(LockHeldError):  # a peer writer is shut out mid-merge
            transaction._acquire_lock(wb.parent / "_state" / "excel.lock")
        return [*existing, _robots("HF-001")]

    result = transaction.merge(wb, "robots_txt", fn, SLUG)
    assert seen == [["R-001"]]
    assert result.rows_affected == 2
    assert _column(wb, "robots_txt") == ["R-001", "HF-001"]


def test_merge_refuses_to_empty_a_populated_sheet(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    committer.commit(wb, "robots_txt", [_robots("R-001")], run_id=RUN_ID, project_slug=SLUG)
    with pytest.raises(EmptyMergeError):
        transaction.merge(wb, "robots_txt", lambda existing: [], SLUG)
    assert _column(wb, "robots_txt") == ["R-001"]


def test_merge_validates_fn_output_before_writing(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    committer.commit(wb, "robots_txt", [_robots("R-001")], run_id=RUN_ID, project_slug=SLUG)
    bad = {**_robots("HF-001"), "level": "NOT-A-LEVEL"}
    with pytest.raises(transaction.RowSchemaError):
        transaction.merge(wb, "robots_txt", lambda existing: [*existing, bad], SLUG)
    assert _column(wb, "robots_txt") == ["R-001"]


# --- two skills merging concurrently ----------------------------------------

def _skill(wb: str, prefix: str, rounds: int) -> None:
    for i in range(1, rounds + 1):
        sheet_merge.merge_prefixed_rows(
            Path(wb), "robots_txt", [_robots("x") for _ in range(i % 3 + 1)],
            id_prefix=prefix, run_id=RUN_ID, project_slug=SLUG, writer=prefix.lower(),
        )


def test_two_concurrent_merges_keep_both_namespaces(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    committer.commit(wb, "robots_txt", [_robots("R-001")], run_id=RUN_ID, project_slug=SLUG)
    ctx = mp.get_context("fork")
    rounds = 6
    procs = [ctx.Process(target=_skill, args=(str(wb), p, rounds)) for p in ("HF-", "RP-")]
    for p in procs:
        p.start()
    for p in procs:
        p.join(timeout=120)
        assert p.exitcode == 0
    ids = _column(wb, "robots_txt")
    expected = rounds % 3 + 1  # each namespace's LAST merge wins inside its own namespace
    assert ids.count("R-001") == 1
    assert sorted(i for i in ids if i.startswith("HF-")) == [f"HF-{n:03d}" for n in range(1, expected + 1)]
    assert sorted(i for i in ids if i.startswith("RP-")) == [f"RP-{n:03d}" for n in range(1, expected + 1)]


# --- sf_import keeps other writers' rows (V-04) -----------------------------

def test_sf_import_keeps_foreign_robots_and_migration_rows(tmp_path, monkeypatch) -> None:
    wb = _wb(tmp_path)
    committer.commit(wb, "robots_txt", [_robots("R-001"), _robots("HF-001"), _robots("RP-001")],
                     run_id=RUN_ID, project_slug=SLUG)
    committer.commit(wb, "redirect_404", [
        _redirect("https://x.test/old-sf", "4xx broken — fix or 301 to a valid URL; linked 1x internally"),
        _redirect("https://x.test/moved", "301"),
        _redirect("https://x.test/gone", "410"),
    ], run_id=RUN_ID, project_slug=SLUG)

    fresh = {
        "robots_txt": [_robots("R-001"), _robots("R-002")],
        "redirect_404": [
            _redirect("https://x.test/new-sf", "3xx — point internal links directly to /t"),
            _redirect("https://x.test/moved", "3xx — point internal links directly to /t"),
        ],
    }
    monkeypatch.setattr(sf_projection, "project_all", lambda raw_dir: fresh)
    sf_import.project_and_write(tmp_path / "raw", wb, SLUG, workspace_root=tmp_path)

    assert _column(wb, "robots_txt") == ["HF-001", "RP-001", "R-001", "R-002"]
    by_url = dict(zip(_column(wb, "redirect_404", 0), _column(wb, "redirect_404", 2)))
    assert by_url == {
        "https://x.test/moved": "301",   # migration plan kept; SF duplicate not added
        "https://x.test/gone": "410",
        "https://x.test/new-sf": "3xx — point internal links directly to /t",
    }  # old-sf (the previous SF snapshot) is refreshed away

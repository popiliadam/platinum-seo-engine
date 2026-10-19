"""Independent-review follow-ups to transaction.merge / sheet_merge (wave 1).

6. Foreign rows read back from the sheet are normalised (an Excel-typed cell
   must not make the whole merge fail validation) — same rule as update() (V-10).
7. A re-run that legitimately finds nothing may clear the caller's OWN rows when
   they are all that the sheet holds; emptying other writers' rows stays refused.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.excel import transaction
from scripts.excel.transaction import EmptyMergeError
from scripts.orchestration import committer
from scripts.util import sheet_merge

SLUG = "test-proj"
RUN = "test-proj-2026-10-19-ab12"


def _wb(tmp_path: Path) -> Path:
    proj = tmp_path / "projects" / SLUG
    (proj / "_state").mkdir(parents=True)
    return proj / "master.xlsx"


def _robots(id_: str) -> dict:
    return {"id": id_, "level": "MEDIUM", "issue": "i", "detail": "d", "resolution": "r"}


def _ids(wb: Path) -> list:
    ws = load_workbook(wb)["robots_txt"]
    return [r[0] for r in ws.iter_rows(min_row=5, values_only=True) if r and r[0] is not None]


def test_zero_findings_rerun_clears_only_own_rows(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    sheet_merge.merge_prefixed_rows(wb, "robots_txt", [_robots("x"), _robots("y")],
                                    id_prefix="HF-", run_id=RUN, project_slug=SLUG, writer="hf")
    sheet_merge.merge_prefixed_rows(wb, "robots_txt", [], id_prefix="HF-",
                                    run_id=RUN, project_slug=SLUG, writer="hf")
    assert _ids(wb) == []


def test_empty_merge_that_would_drop_foreign_rows_is_still_refused(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    committer.commit(wb, "robots_txt", [_robots("R-001")], run_id=RUN, project_slug=SLUG)
    with pytest.raises(EmptyMergeError):
        transaction.merge(wb, "robots_txt", lambda existing: [], SLUG)
    assert _ids(wb) == ["R-001"]


def test_foreign_excel_typed_cell_does_not_block_the_merge(tmp_path: Path) -> None:
    wb = _wb(tmp_path)
    rows = [{"url": "/a", "inlinks": 3, "action": "301", "target_url": "/b", "status": "TODO"}]
    committer.commit(wb, "redirect_404", rows, run_id=RUN, project_slug=SLUG)
    book = load_workbook(wb)
    book["redirect_404"]["B5"] = "3"  # a number typed as text in Excel
    book.save(wb)
    new = [{"url": "/c", "inlinks": 1, "action": "410", "target_url": "", "status": "TODO"}]
    sheet_merge.merge_keyed_rows(wb, "redirect_404", new, run_id=RUN, project_slug=SLUG,
                                 writer="migration-map")
    ws = load_workbook(wb)["redirect_404"]
    assert [ws["A5"].value, ws["A6"].value] == ["/a", "/c"]

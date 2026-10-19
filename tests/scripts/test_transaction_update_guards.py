"""V-10 / V-11 / V-12 — transaction.update() and the "workbook is open" guard.

V-10: Excel turns a typed date into a real datetime cell; update() re-validated
      the whole merged row with raw cell values, so editing ANY other column of
      that row failed (live: dentnotion 1, rkturizm-tr 2, lastiksa-tr 1 cells).
V-11: an empty where={} — or a where key that is not a column (always None) —
      matched every row, so one call could rewrite a whole sheet.
V-12: only Excel's ~$ owner file was detected; LibreOffice writes .~lock.<name>#.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.excel import transaction
from scripts.excel.transaction import RowSchemaError, WorkbookLockedByExcelError

SLUG = "test-proj"


def _wb(tmp_path: Path) -> Path:
    proj = tmp_path / "projects" / SLUG
    (proj / "_state").mkdir(parents=True)
    return proj / "master.xlsx"


def _row(i: int) -> dict:
    return {"id": f"CW-{i}", "task_or_content": "t", "url": f"/u{i}",
            "date": "2026-10-01", "category": "tech", "note": ""}


def _seed(wb_path: Path, n: int = 2) -> None:
    transaction.append(wb_path, "completed_work", [_row(i) for i in range(1, n + 1)], SLUG)


def test_update_works_on_a_row_whose_date_excel_converted(tmp_path: Path) -> None:
    wb_path = _wb(tmp_path)
    _seed(wb_path)
    wb = load_workbook(wb_path)
    ws = wb["completed_work"]
    ws["D5"] = datetime(2026, 10, 1)  # what Excel does to a typed date
    wb.save(wb_path)

    result = transaction.update(wb_path, "completed_work", {"id": "CW-1"},
                                {"note": "verified"}, SLUG)
    assert result.rows_affected == 1
    ws = load_workbook(wb_path)["completed_work"]
    assert ws["F5"].value == "verified"


def test_update_with_empty_where_is_refused(tmp_path: Path) -> None:
    wb_path = _wb(tmp_path)
    _seed(wb_path)
    with pytest.raises(RowSchemaError, match="where"):
        transaction.update(wb_path, "completed_work", {}, {"note": "all"}, SLUG)
    notes = [c.value for c in load_workbook(wb_path)["completed_work"]["F"][4:]]
    assert "all" not in notes


def test_update_with_unknown_where_column_is_refused(tmp_path: Path) -> None:
    wb_path = _wb(tmp_path)
    _seed(wb_path)
    with pytest.raises(RowSchemaError, match="tsk_id"):
        transaction.update(wb_path, "completed_work", {"tsk_id": None}, {"note": "all"}, SLUG)


def test_libreoffice_lock_blocks_the_write(tmp_path: Path) -> None:
    wb_path = _wb(tmp_path)
    _seed(wb_path, 1)
    (wb_path.parent / f".~lock.{wb_path.name}#").write_text("user,host", encoding="utf-8")
    with pytest.raises(WorkbookLockedByExcelError, match="lock"):
        transaction.append(wb_path, "completed_work", [_row(9)], SLUG)

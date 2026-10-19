"""K-01 / K-03 — CI step 1 runs the real invariant validator and can go red.

The old step executed SKILL.md blocks whose checks were commented out and whose
result was a hard-coded GREEN. drift_gate runs validate_invariants on a copy of
the template workbook; a broken invariant must turn it red (judge self-test).
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import yaml
from openpyxl import load_workbook

from scripts.ci import drift_gate as dg

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "master-excel.xlsx"


def _header_row(sheet: str) -> int:
    schema = json.loads((ROOT / "schemas" / "master-excel.schema.json").read_text(encoding="utf-8"))
    return int(schema["sheets"][sheet]["header_row"])


def test_template_passes_with_the_known_unimplemented_rules() -> None:
    rc, report = dg.run_gate(TEMPLATE)
    assert rc == 0, report
    assert report["overall"] == "AMBER" and report["evaluated"] >= dg.MIN_EVALUATED
    assert set(report["not_implemented"]) == dg.EXPECTED_NOT_IMPLEMENTED


def test_duplicate_task_id_turns_the_gate_red(tmp_path: Path) -> None:
    broken = tmp_path / "fixture-broken.xlsx"
    shutil.copyfile(TEMPLATE, broken)
    wb = load_workbook(broken)
    ws = wb["master_task"]
    hr = _header_row("master_task")
    headers = [c.value for c in ws[hr]]
    col = headers.index("task_id") + 1
    for offset in (1, 2):  # the same task_id twice -> F-09 FAIL
        ws.cell(row=hr + offset, column=col, value="T-00001")
    wb.save(broken)
    rc, report = dg.run_gate(broken)
    assert rc == 1 and report["overall"] == "RED", report
    assert any("F-09" in p for p in report["problems"]), report


def test_ci_step_1_runs_the_gate() -> None:
    cfg = yaml.safe_load((ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8"))
    step = next(s for s in cfg["jobs"]["ci"]["steps"] if s.get("name") == "1. drift-check")
    assert step["run"].strip() == "python3 scripts/ci/drift_gate.py"

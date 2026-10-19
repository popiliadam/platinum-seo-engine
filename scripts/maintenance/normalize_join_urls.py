#!/usr/bin/env python3
"""normalize_join_urls.py — one-time repair for D-03 drift in engine join sheets.

Rows written before the D-05 / D-07 fixes carry raw URLs (trailing slash,
upper-case host, tracking params) in sheets the engine joins on, so
``validate_invariants`` D-03 reports FAIL (routed to manual triage). This
script rewrites those cells to ``normalize_url`` form. Nothing else changes.

Dry-run by default (prints what would change). ``--apply`` writes each sheet
through ``transaction.merge`` (excel.lock, backup, atomic save, provenance).

    python3 scripts/maintenance/normalize_join_urls.py --project <slug> [--apply]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from scripts.excel import transaction  # noqa: E402
from scripts.util.url_normalize import normalize_url  # noqa: E402
from scripts.validation import validate_invariants as vi  # noqa: E402


def _norm(value):
    if not isinstance(value, str) or not value.lower().startswith(("http://", "https://")):
        return value
    try:
        return normalize_url(value)
    except ValueError:
        return value


def normalized_rows(rows: list[dict], url_cols: list[str]) -> tuple[list[dict], int]:
    """New row dicts with URL columns normalized, and how many cells changed."""
    out, changed = [], 0
    for row in rows:
        new = dict(row)
        for c in url_cols:
            fixed = _norm(new.get(c))
            if fixed != new.get(c):
                new[c] = fixed
                changed += 1
        out.append(new)
    return out, changed


def _url_cols(sheet: str) -> list[str]:
    schema = vi._load_master_excel_schema()
    cols = [c.get("name") for c in schema["sheets"][sheet].get("required_columns") or []]
    return [c for c in cols if c in vi._URL_COLUMN_NAMES]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="normalize_join_urls.py")
    ap.add_argument("--project", required=True)
    ap.add_argument("--workspace-root", default=None)
    ap.add_argument("--apply", action="store_true", help="write (default: dry-run)")
    args = ap.parse_args(argv)
    ws = vi._resolve_workspace_root(Path(args.workspace_root) if args.workspace_root else None)
    wb_path = vi._project_dir(args.project, ws) / "master.xlsx"
    if not wb_path.is_file():
        print(f"master.xlsx not found: {wb_path}", file=sys.stderr)
        return 2
    from openpyxl import load_workbook
    wb = load_workbook(str(wb_path), read_only=True, data_only=True)
    try:
        plan = {}
        for sheet in sorted(vi._D03_JOIN_SHEETS):
            if sheet not in wb.sheetnames:
                continue
            cols = _url_cols(sheet)
            _rows, changed = normalized_rows(vi._iter_rows_as_dicts(wb, sheet), cols)
            if changed:
                plan[sheet] = (cols, changed)
    finally:
        wb.close()
    for sheet, (cols, changed) in plan.items():
        print(f"{sheet}: {changed} URL cell(s) to normalize ({', '.join(cols)})")
    if not plan:
        print("nothing to normalize")
        return 0
    if not args.apply:
        print("dry-run — re-run with --apply to write")
        return 0
    for sheet, (cols, _changed) in plan.items():
        transaction.merge(wb_path, sheet, lambda rows, c=cols: normalized_rows(rows, c)[0],
                          args.project, writer="human")
        print(f"{sheet}: written")
    return 0


if __name__ == "__main__":
    sys.exit(main())
